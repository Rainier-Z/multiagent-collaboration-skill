"""End-to-end regression coverage for coordinator binding and delivery gates."""

from __future__ import annotations

import json
import base64
import hashlib
import os
import shutil
import sys
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import patch
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from orchestrate_discussion import FakeWakeAdapter, orchestrate_once
from participant_runtime.protocol import (
    Instruction,
    isolation_evidence_path,
    load_isolation_evidence,
    load_stop_attestation,
    stop_attestation_path,
)
from workflow_core import WorkflowError, sha256_file
import confirm_decision
import export_docx
from docx import Document
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from attestation_keys import sign_payload


def fail_open(_path):
    raise OSError("simulated opener failure")


def write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


class EndToEndV2(unittest.TestCase):
    def setUp(self) -> None:
        self.workspace = Path(tempfile.mkdtemp(prefix="multiagent-e2e-v2-"))
        self.key_dir = tempfile.TemporaryDirectory(prefix="multiagent-e2e-key-")
        self.key_id = "e2e-test-key"
        self.private_key_path = Path(self.key_dir.name) / "private.pem"
        private_key = Ed25519PrivateKey.generate()
        self.private_key_path.write_bytes(private_key.private_bytes(
            encoding=serialization.Encoding.PEM,
            format=serialization.PrivateFormat.PKCS8,
            encryption_algorithm=serialization.NoEncryption(),
        ))
        self.public_key_b64 = base64.b64encode(private_key.public_key().public_bytes(
            encoding=serialization.Encoding.Raw,
            format=serialization.PublicFormat.Raw,
        )).decode("ascii")
        self.attestation_env = patch.dict(os.environ, {
            "MULTIAGENT_ATTESTATION_KEY_ID": self.key_id,
            "MULTIAGENT_ATTESTATION_PUBLIC_KEY_B64": self.public_key_b64,
            "MULTIAGENT_ATTESTATION_TRUSTED_KEYS_JSON": json.dumps({self.key_id: self.public_key_b64}),
        })
        self.attestation_env.start()
        self.internal = self.workspace / ".multiagent"
        self.internal.mkdir()
        self.participants = ["planner-ds", "claude-a", "openclaw"]
        self.bindings = {
            "planner-ds": {"platform_id": "deepseek", "session_id": "planner-session-1"},
            "claude-a": {"platform_id": "claude-code", "session_id": "claude-session-7"},
            "openclaw": {"platform_id": "openclaw", "session_id": "openclaw-session-3"},
        }
        self.discussion = self.workspace / "测试讨论文档_2026-09-13.md"
        self.discussion.write_text(
            "# 测试讨论\n\n## 二、独立提案\n\n## 三、交叉回应\n\n"
            "## 四、结构化决策包\n\n## 五、候选决策与 Word 审阅\n\n"
            "## 六、确认固化记录\n\n## 七、正式 Word 交付\n",
            encoding="utf-8",
        )
        (self.workspace / "project-context.md").write_text("# 目标\n完成独立讨论与决策。\n", encoding="utf-8")
        self.state = {
            "protocol_version": "1.0",
            "discussion_id": "e2e-test",
            "stage": "initialized",
            "expected_participants": self.participants,
            "submission_status": {a: "pending" for a in self.participants},
            "response_status": {a: "pending" for a in self.participants},
            "coordinator": "planner-ds",
            "coordinator_binding": {
                "agent_id": "planner-ds", "role": "coordinator",
                "platform_id": "deepseek", "session_id": "planner-session-1",
            },
            "participant_bindings": self.bindings,
            "candidate_decision_ids": [],
            "confirmed_decision_ids": [],
            "proposal_disposition": "delete",
            "coordination_lease_until": "2026-09-13T13:00:00+08:00",
            "coordinator_timeout": 300,
            "participant_timeout": 900,
            "last_checked_at": "2026-09-13T12:00:00+08:00",
            "monitoring": {"enabled": True, "mode": "reply_before", "status": "active"},
            "revision": 1,
            "runtime_distribution": {"version": "1.0.0"},
            "retry_policy": {"max_attempts": 3},
            "content_authority": {
                "discussion_path": self.discussion.name,
                "sha256": sha256_file(self.discussion),
            },
        }
        write_json(self.internal / "state.json", self.state)
        for index, agent in enumerate(self.participants, 1):
            write_json(self.internal / "receipts" / agent / f"B-{index}-completed.json", {
                "instruction_id": f"B-{index}", "agent_id": agent,
                "status": "completed", "kind": "bootstrap",
            })

    def tearDown(self) -> None:
        self.attestation_env.stop()
        self.key_dir.cleanup()
        shutil.rmtree(self.workspace, ignore_errors=True)

    def read_state(self) -> dict:
        return json.loads((self.internal / "state.json").read_text(encoding="utf-8"))

    def run_as_coordinator(self, *, opener=None, open_candidate=True):
        return orchestrate_once(
            self.workspace, FakeWakeAdapter(), open_candidate=open_candidate,
            actor="planner-ds", platform_id="deepseek", session_id="planner-session-1",
            opener=opener,
        )

    def instructions(self, agent: str, kind: str) -> list[Instruction]:
        return [Instruction.from_dict(json.loads(path.read_text(encoding="utf-8")))
                for path in sorted((self.internal / "instructions" / agent).glob(f"*-{kind}-*.json"))]

    def attest_proposal(self, instruction: Instruction) -> dict:
        scope = instruction.access_scope
        attestation = {
            "issuer": "platform-attestation", "source": "platform-attestation",
            "evidence_type": "platform_sandbox", "enforcement_type": "platform_sandbox",
            "issued_at": "2026-09-13T12:00:00+08:00",
            "agent_id": instruction.agent_id, "instruction_id": instruction.instruction_id,
            "platform_id": instruction.platform_id, "session_id": instruction.session_id,
            "view_root": scope["view_root"], "scope_digest": scope["scope_digest"],
            "input_manifest_sha256": scope["input_manifest_sha256"],
            "allowed_read_roots": scope["allowed_read_roots"],
            "allowed_write_roots": scope["allowed_write_roots"],
            "trust": {"trust_root_id": "test-platform-fixture"},
        }
        attestation = sign_payload(
            attestation, private_key_path=self.private_key_path, key_id=self.key_id,
        )
        path = isolation_evidence_path(self.workspace, instruction)
        write_json(path, attestation)
        return load_isolation_evidence(self.workspace, instruction)

    def complete_outputs(self, agent: str, kind: str, text: str, *, evidence=None) -> None:
        instruction = self.instructions(agent, kind)[-1]
        output = self.workspace / instruction.output_path
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text, encoding="utf-8")
        payload = {
            "instruction_id": instruction.instruction_id, "agent_id": agent,
            "kind": kind, "status": "completed",
            "output_path": instruction.output_path, "output_sha256": sha256_file(output),
        }
        if kind in {"propose", "repair"}:
            payload["isolation_evidence"] = evidence or self.attest_proposal(instruction)
        else:
            payload["isolation_evidence"] = {"mode": "restricted_view", "contract_validation": "passed"}
        write_json(self.internal / "receipts" / agent / f"{instruction.instruction_id}-completed.json", payload)

    def complete_stop(self, agent: str) -> None:
        instruction = self.instructions(agent, "stop")[-1]
        now = datetime.now(timezone.utc).astimezone()
        issued_at = datetime.fromisoformat(instruction.issued_at)
        stopped_at = max(now - timedelta(seconds=10), issued_at + timedelta(milliseconds=1))
        marker = self.internal / "receipts" / agent / f"{instruction.instruction_id}-stop-marker.txt"
        marker.parent.mkdir(parents=True, exist_ok=True)
        marker.write_text("stopped\n", encoding="utf-8")
        attestation = {
            "discussion_id": "e2e-test",
            "agent_id": agent,
            "instruction_id": instruction.instruction_id,
            "instruction_sha256": instruction.sha256,
            "platform_id": instruction.platform_id,
            "session_id": instruction.session_id,
            "stopped_at": stopped_at.isoformat(timespec="milliseconds"),
            "mechanism": "session_monitor_stop",
            "action_verified": True,
            "target": "current session monitor instance",
            "proof_reference": ".multiagent/audit/platform-evidence/proofs/%s/%s.json" % (
                agent, instruction.instruction_id,
            ),
        }
        if instruction.platform_id == "openclaw":
            attestation.update({
                "automation_job_id": "openclaw-job-1",
                "removal_verified": True,
                "removal_checked_at": (stopped_at + timedelta(milliseconds=1)).isoformat(timespec="milliseconds"),
            })
        proof_artifact = {
            field: attestation[field]
            for field in (
                "discussion_id", "instruction_id", "instruction_sha256", "agent_id", "platform_id", "session_id",
                "stopped_at", "mechanism", "target", "action_verified",
            )
        }
        if instruction.platform_id == "openclaw":
            proof_artifact.update({field: attestation[field] for field in (
                "automation_job_id", "removal_verified", "removal_checked_at",
            )})
        proof_path = self.workspace / attestation["proof_reference"]
        write_json(proof_path, proof_artifact)
        attestation["proof_sha256"] = hashlib.sha256(proof_path.read_bytes()).hexdigest()
        signed_attestation = sign_payload(
            attestation, private_key_path=self.private_key_path, key_id=self.key_id,
        )
        write_json(stop_attestation_path(self.workspace, instruction), signed_attestation)
        evidence = load_stop_attestation(self.workspace, instruction)
        receipt = {
            "instruction_id": instruction.instruction_id,
            "agent_id": agent,
            "kind": "stop",
            "status": "completed",
            "runtime_version": instruction.runtime_version,
            "state_revision": instruction.state_revision,
            "at": datetime.now(timezone.utc).astimezone().isoformat(timespec="milliseconds"),
            "attempt": instruction.attempt,
            "output_path": marker.relative_to(self.workspace).as_posix(),
            "output_sha256": sha256_file(marker),
            "isolation_evidence": evidence,
        }
        write_json(self.internal / "receipts" / agent / f"{instruction.instruction_id}-completed.json", receipt)

    def drive_to_candidate(self) -> None:
        self.run_as_coordinator()
        for agent in self.participants:
            instruction = self.instructions(agent, "propose")[-1]
            self.complete_outputs(agent, "propose", f"# {agent} proposal\n观点独立。\n")
        result = self.run_as_coordinator()
        self.assertEqual(result.stage, "cross_response")
        for agent in self.participants:
            instruction = self.instructions(agent, "respond")[-1]
            self.assertTrue(any(path.endswith("/discussion.md") for path in instruction.input_paths))
            self.complete_outputs(agent, "respond", f"# {agent} response\n回应共同讨论。\n")
        self.assertEqual(self.run_as_coordinator(opener=lambda _path: False).stage, "candidate_decision")

    def test_non_claude_coordinator_must_match_exact_session(self) -> None:
        before = (self.internal / "state.json").read_bytes()
        with self.assertRaises(WorkflowError):
            orchestrate_once(self.workspace, FakeWakeAdapter(), actor="planner-ds",
                             platform_id="deepseek", session_id="wrong-session")
        self.assertEqual((self.internal / "state.json").read_bytes(), before)
        self.assertEqual(self.run_as_coordinator().stage, "independent_proposal")
        for agent in self.participants:
            instruction = self.instructions(agent, "propose")[-1]
            self.assertEqual(instruction.platform_id, self.bindings[agent]["platform_id"])
            self.assertEqual(instruction.session_id, self.bindings[agent]["session_id"])

    def test_missing_participant_binding_fails_closed(self) -> None:
        self.state["participant_bindings"].pop("claude-a")
        write_json(self.internal / "state.json", self.state)
        before = (self.internal / "state.json").read_bytes()
        with self.assertRaises(WorkflowError):
            self.run_as_coordinator()
        self.assertEqual((self.internal / "state.json").read_bytes(), before)
        self.assertFalse(list((self.internal / "instructions").rglob("*-propose-*.json")))

    def test_missing_or_forged_attestation_does_not_merge(self) -> None:
        self.run_as_coordinator()
        for agent in self.participants:
            self.complete_outputs(agent, "propose", "# proposal\n", evidence={"authenticity": "self-reported"})
        result = self.run_as_coordinator()
        self.assertEqual(result.stage, "independent_proposal")
        self.assertNotIn("cross_response", self.read_state()["stage"])
        self.assertNotIn("# proposal", self.discussion.read_text(encoding="utf-8"))

    def test_qualified_attestation_merges_and_candidate_open_gates_confirmation(self) -> None:
        self.run_as_coordinator()
        for agent in self.participants:
            self.complete_outputs(agent, "propose", f"# {agent} proposal\n提案内容。\n")
        self.assertEqual(self.run_as_coordinator().stage, "cross_response")
        merged = self.discussion.read_text(encoding="utf-8")
        self.assertIn("## 二、独立提案", merged)
        self.assertIn("planner-ds proposal", merged)
        self.assertIn("## 三、交叉回应", merged)
        for agent in self.participants:
            self.complete_outputs(agent, "respond", f"# {agent} response\n回应内容。\n")
        self.assertEqual(self.run_as_coordinator(open_candidate=False).stage, "candidate_decision")
        failed = self.run_as_coordinator(opener=fail_open)
        self.assertEqual(failed.stage, "candidate_decision")
        state = self.read_state()
        self.assertFalse(state["candidate_delivery"]["opened"])
        self.assertTrue(state["monitoring"]["enabled"])
        success = self.run_as_coordinator(opener=lambda _path: True)
        self.assertEqual(success.stage, "candidate_decision")
        state = self.read_state()
        self.assertTrue(state["monitoring"]["enabled"])
        self.assertEqual(state["monitoring"]["status"], "stopping")
        stop_instructions = {
            agent: self.instructions(agent, "stop") for agent in self.participants
        }
        self.assertTrue(all(len(items) == 1 for items in stop_instructions.values()))
        self.assertEqual(stop_instructions["openclaw"][0].platform_id, "openclaw")
        self.assertEqual(stop_instructions["openclaw"][0].session_id, "openclaw-session-3")
        self.complete_stop(self.participants[0])
        waiting = self.run_as_coordinator()
        self.assertEqual(waiting.stage, "candidate_decision")
        self.assertTrue(self.read_state()["monitoring"]["enabled"])
        for agent in self.participants[1:]:
            self.complete_stop(agent)
        stopped = self.run_as_coordinator()
        self.assertEqual(stopped.stage, "user_confirmation")
        state = self.read_state()
        self.assertFalse(state["monitoring"]["enabled"])
        self.assertEqual(state["monitoring"]["status"], "stopped")
        self.assertTrue(state["monitoring"]["stop_requested_at"])
        self.assertTrue(state["monitoring"]["stopped_at"])
        self.assertEqual(len(list((self.internal / "instructions").rglob("*-stop-*.json"))), len(self.participants))
        self.assertTrue((self.workspace / "候选决策.docx").is_file())
        candidate_doc = "\n".join(p.text for p in Document(self.workspace / "候选决策.docx").paragraphs)
        self.assertIn("planner-ds proposal", candidate_doc)
        self.assertIn("claude-a response", candidate_doc)
        confirm_argv = ["confirm_decision.py", "--state", str(self.internal / "state.json"),
                        "--actor", "planner-ds", "--candidate-id", "C-0001",
                        "--platform-id", "deepseek", "--session-id", "planner-session-1",
                        "--confirm-text", "确认 C-0001"]
        with patch("sys.argv", confirm_argv), patch.object(export_docx, "open_document", return_value=False):
            self.assertNotEqual(confirm_decision.main(), 0)
        self.assertEqual(self.read_state()["stage"], "confirmed_decision")
        with patch("sys.argv", confirm_argv), patch.object(export_docx, "open_document", return_value=True):
            self.assertEqual(confirm_decision.main(), 0)
        self.assertEqual(self.read_state()["stage"], "delivered")
        self.assertTrue((self.workspace / "最终决策.docx").is_file())
        final_doc = "\n".join(p.text for p in Document(self.workspace / "最终决策.docx").paragraphs)
        self.assertIn("planner-ds proposal", final_doc)
        self.assertIn("claude-a response", final_doc)
        roots = {p.name for p in self.workspace.iterdir() if p.is_file()}
        self.assertEqual(roots, {"project-context.md", self.discussion.name, "候选决策.docx", "最终决策.docx"})


if __name__ == "__main__":
    unittest.main()
